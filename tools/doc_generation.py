from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


def generate_docx(sections: dict, out_path: str) -> str:
    """
    Create a formal approval-note Word document.

    Expected sections format:
    {
        "title": "Approval Note",
        "body": {
            "Background": "....",
            "Observations": "....",
            "Non-Conformities": "....",
            "Recommendation": "...."
        }
    }
    """
    output_path = Path(out_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(11)

    title = doc.add_heading(sections.get("title", "Approval Note"), level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for heading, body in sections.get("body", {}).items():
        doc.add_heading(heading, level=2)
        paragraph = doc.add_paragraph(str(body))
        paragraph.paragraph_format.space_after = Pt(8)

    doc.add_paragraph()
    footer = doc.add_paragraph("Generated locally by Sovereign AI Workbench")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(9)

    doc.save(output_path)
    return str(output_path)


if __name__ == "__main__":
    sample_sections = {
        "title": "Approval Note",
        "body": {
            "Background": (
                "A routine boiler pressure inspection was conducted for "
                "Boiler Unit B-07."
            ),
            "Observations": (
                "The recorded operating pressure was 13.2 bar. "
                "The inspection record was reviewed."
            ),
            "Non-Conformities": (
                "The reading exceeds the 12 bar routine-operation limit "
                "specified under SOP-101, Clause 4.2."
            ),
            "Recommendation": (
                "Escalate the deviation to the shift engineer and carry out "
                "a pressure-control system inspection before further operation."
            ),
        },
    }

    created_file = generate_docx(
        sections=sample_sections,
        out_path="data/uploads/sample_approval_note.docx",
    )

    print(f"Document created: {created_file}")