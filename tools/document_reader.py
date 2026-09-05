import os
from pathlib import Path

import pymupdf
import pandas as pd
from docx import Document


SUPPORTED_DOCUMENT_TYPES = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".csv",
    ".xlsx",
    ".xls",
}


def get_file_metadata(file_path: str) -> dict:
    """
    Return safe local metadata about an uploaded document.

    This function does not send the file anywhere. It only reads local
    filesystem metadata needed for routing and user-facing messages.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    return {
        "filename": path.name,
        "extension": path.suffix.lower(),
        "size_bytes": path.stat().st_size,
        "size_mb": round(path.stat().st_size / (1024 * 1024), 2),
    }


def read_text_file(file_path: str) -> dict:
    """Read a local plain-text or Markdown document."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as file:
        text = file.read()

    return {
        "content": text,
        "metadata": {
            "format": "text",
            "characters": len(text),
        },
    }


def read_pdf(file_path: str, max_pages: int = 30) -> dict:
    """
    Extract selectable text from a local PDF using PyMuPDF.

    If a scanned PDF has little/no selectable text, this function reports
    that fact. Vision/OCR fallback will be added later.
    """
    pdf = pymupdf.open(file_path)
    page_count = len(pdf)
    pages_to_read = min(page_count, max_pages)

    page_texts = []

    for page_index in range(pages_to_read):
        page = pdf[page_index]
        text = page.get_text("text").strip()

        if text:
            page_texts.append(
                f"\n--- Page {page_index + 1} ---\n{text}"
            )

    pdf.close()

    content = "\n".join(page_texts).strip()
    is_likely_scanned = len(content) < 100 and page_count > 0

    return {
        "content": content,
        "metadata": {
            "format": "pdf",
            "page_count": page_count,
            "pages_read": pages_to_read,
            "is_likely_scanned": is_likely_scanned,
            "characters_extracted": len(content),
        },
    }


def read_docx(file_path: str) -> dict:
    """Extract paragraphs and simple tables from a local DOCX file."""
    document = Document(file_path)

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    table_sections = []

    for table_index, table in enumerate(document.tables, start=1):
        rows = []

        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]
            rows.append(" | ".join(cells))

        if rows:
            table_sections.append(
                f"\n--- Table {table_index} ---\n" + "\n".join(rows)
            )

    content_parts = paragraphs + table_sections
    content = "\n".join(content_parts).strip()

    return {
        "content": content,
        "metadata": {
            "format": "docx",
            "paragraph_count": len(paragraphs),
            "table_count": len(table_sections),
            "characters_extracted": len(content),
        },
    }


def read_csv(file_path: str, max_rows: int = 30) -> dict:
    """
    Read a local CSV and return a compact Markdown-friendly data preview.

    We do not pass an entire massive spreadsheet directly to the LLM.
    This preview is enough for initial summarization and routing.
    """
    dataframe = pd.read_csv(file_path)

    preview = dataframe.head(max_rows)
    content = preview.to_markdown(index=False)

    return {
        "content": content,
        "metadata": {
            "format": "csv",
            "row_count": int(len(dataframe)),
            "column_count": int(len(dataframe.columns)),
            "columns": dataframe.columns.tolist(),
            "preview_rows": int(len(preview)),
        },
    }


def read_excel(file_path: str, max_rows: int = 30) -> dict:
    """
    Read the first sheet of a local Excel workbook and return a preview.

    A later spreadsheet-analysis tool can handle multi-sheet operations,
    formulas, charts, and output workbook creation.
    """
    workbook = pd.ExcelFile(file_path)
    sheet_names = workbook.sheet_names

    first_sheet = sheet_names[0]
    dataframe = pd.read_excel(file_path, sheet_name=first_sheet)

    preview = dataframe.head(max_rows)
    content = preview.to_markdown(index=False)

    return {
        "content": content,
        "metadata": {
            "format": "excel",
            "sheet_names": sheet_names,
            "selected_sheet": first_sheet,
            "row_count": int(len(dataframe)),
            "column_count": int(len(dataframe.columns)),
            "columns": dataframe.columns.tolist(),
            "preview_rows": int(len(preview)),
        },
    }


def read_document(file_path: str) -> dict:
    """
    Universal local document reader.

    Returns:
    {
        "content": "...",
        "metadata": {...}
    }
    """
    metadata = get_file_metadata(file_path)
    extension = metadata["extension"]

    if extension not in SUPPORTED_DOCUMENT_TYPES:
        raise ValueError(
            f"Unsupported document type: {extension or 'no extension'}"
        )

    if extension in {".txt", ".md"}:
        result = read_text_file(file_path)

    elif extension == ".pdf":
        result = read_pdf(file_path)

    elif extension == ".docx":
        result = read_docx(file_path)

    elif extension == ".csv":
        result = read_csv(file_path)

    elif extension in {".xlsx", ".xls"}:
        result = read_excel(file_path)

    else:
        raise ValueError(f"Unsupported document type: {extension}")

    result["metadata"] = {
        **metadata,
        **result["metadata"],
    }

    return result


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print(
            "Usage: python tools/document_reader.py <path-to-document>"
        )
        raise SystemExit(1)

    path = sys.argv[1]

    try:
        result = read_document(path)

        print("\n--- METADATA ---")
        for key, value in result["metadata"].items():
            print(f"{key}: {value}")

        print("\n--- EXTRACTED CONTENT PREVIEW ---")
        print(result["content"][:2000])

    except Exception as error:
        print(f"Document read failed: {error}")
        raise SystemExit(1)